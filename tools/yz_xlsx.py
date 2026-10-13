"""עריכת מאגר דברי הימים בקובץ אקסל.
   ייצוא:  python3 tools/yz_xlsx.py export   → data-src/dvrei-hayamim.xlsx
   ייבוא:  python3 tools/yz_xlsx.py import   → מעדכן את data-src/yz-data.js ובונה מחדש את assets/yz.js
   עמודת "מזהה" מקשרת כל שורה לרשומה באתר (גם לעריכות שמשתמשים שמרו). לא משנים אותה.
   שורה חדשה: משאירים את המזהה ריק. מחיקה: מוחקים את השורה (או את התיאור)."""
import sys, json, subprocess, pathlib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC, XLSX = ROOT / 'data-src/yz-data.js', ROOT / 'data-src/dvrei-hayamim.xlsx'
MONTHS = {7:'תשרי',8:'חשון',9:'כסלו',10:'טבת',11:'שבט',12:'אדר / אדר א׳',13:'אדר ב׳',1:'ניסן',2:'אייר',3:'סיון',4:'תמוז',5:'אב',6:'אלול'}
ORDER = [7,8,9,10,11,12,13,1,2,3,4,5,6]
NAME2M = {v:k for k,v in MONTHS.items()}; NAME2M.update({'אדר':12,'אדר א':12,"אדר א'":12,'אדר ב':13,"אדר ב'":13,'חשוון':8,'כסליו':9,'סיוון':3,'מנחם אב':5})
TYPES = {1:'אישיות (יום פטירה)',0:'אירוע'}
SRCS = {0:'עתים לבינה',1:'מאגר נוסף'}

def load():
    code = "const fs=require('fs');const src=fs.readFileSync(process.argv[1],'utf8');const f=new Function(src+';return {YZ,SEL:YZ_DEFAULT_SEL}');process.stdout.write(JSON.stringify(f()))"
    return json.loads(subprocess.run(['node','-e',code,str(SRC)],capture_output=True,text=True,check=True).stdout)

def export():
    o = load(); sel = set(o['SEL'])
    wb = Workbook(); ws = wb.active; ws.title = 'דברי הימים'; ws.sheet_view.rightToLeft = True
    head = ['מזהה','חודש','יום','שנה','תיאור','סוג','מקור','מופיע כברירת מחדל']
    ws.append(head)
    rows = [(i,r) for i,r in enumerate(o['YZ']) if r[3]]
    rows.sort(key=lambda x:(ORDER.index(x[1][0]) if x[1][0] in ORDER else 99, x[1][1], x[0]))
    for i,r in rows:
        ws.append([i, MONTHS.get(r[0],str(r[0])), r[1], r[2], r[3], TYPES[1 if r[4] else 0], SRCS[1 if len(r)>5 and r[5] else 0], 'כן' if i in sel else ''])
    bold = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='1F5C63')
    for c in ws[1]: c.font = bold; c.fill = fill; c.alignment = Alignment(horizontal='center', vertical='center')
    for col,w in zip('ABCDEFGH',[8,13,6,10,90,20,13,12]): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[4].alignment = Alignment(wrap_text=True, vertical='top')
        for c in (row[0],row[1],row[2],row[3],row[5],row[6],row[7]): c.alignment = Alignment(vertical='top', horizontal='center')
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = 'A1:H%d' % ws.max_row
    n = ws.max_row + 500
    for rng, opts in (('B2:B%d'%n, list(MONTHS.values())), ('F2:F%d'%n, list(TYPES.values())), ('H2:H%d'%n, ['כן',''])):
        dv = DataValidation(type='list', formula1='"%s"' % ','.join(o for o in opts if o), allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
    dv = DataValidation(type='whole', operator='between', formula1='1', formula2='30', allow_blank=False); ws.add_data_validation(dv); dv.add('C2:C%d'%n)
    h = wb.create_sheet('הוראות'); h.sheet_view.rightToLeft = True; h.column_dimensions['A'].width = 110
    for line in ['איך עורכים את מאגר דברי הימים',
        '• עורכים כל תא בגליון "דברי הימים": חודש, יום, שנה (בעברית, כמו בקובץ), תיאור, סוג.',
        '• "סוג": אישיות = יום פטירה (יארצייט) של אדם; אירוע = מאורע היסטורי.',
        '• "מופיע כברירת מחדל": "כן" = הרשומה מסומנת מראש בלוח חדש.',
        '• שורה חדשה: מוסיפים בסוף הגליון ומשאירים את "מזהה" ריק.',
        '• מחיקה: מוחקים את כל השורה.',
        '• לא משנים את עמודת "מזהה": היא מקשרת לעריכות שמשתמשים כבר שמרו באתר.',
        '• בסיום: שולחים את הקובץ, והוא נטען לאתר (python3 tools/yz_xlsx.py import).',
        '• אדר בשנה פשוטה = "אדר / אדר א׳". אדר ב׳ = רק בשנה מעוברת.']:
        h.append([line])
    h['A1'].font = Font(bold=True, size=14)
    wb.save(XLSX); print('exported', len(rows), '→', XLSX)

def esc(s): return json.dumps(s, ensure_ascii=False)

def imp():
    o = load(); YZ = [list(r) for r in o['YZ']]; seen = set(); sel = []
    ws = load_workbook(XLSX)['דברי הימים']
    hdr = [c.value for c in ws[1]]; ix = {k:hdr.index(k) for k in ['מזהה','חודש','יום','שנה','תיאור','סוג','מקור','מופיע כברירת מחדל']}
    added = changed = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        g = lambda k: row[ix[k]] if ix[k] < len(row) else None
        txt = str(g('תיאור') or '').strip()
        mname = str(g('חודש') or '').strip(); m = NAME2M.get(mname)
        if not txt or not m: continue
        d = int(g('יום') or 1); y = str(g('שנה') or '').strip(); p = 1 if str(g('סוג') or '').startswith('אישיות') else 0
        srcf = 1 if str(g('מקור') or '') == SRCS[1] else 0
        rec = [m, d, y, txt, p] + ([1] if srcf else [])
        rid = g('מזהה')
        if rid is not None and str(rid).strip() != '' and 0 <= int(rid) < len(YZ):
            i = int(rid)
            if YZ[i][:5] != rec[:5]: changed += 1
            YZ[i] = rec
        else:
            YZ.append(rec); i = len(YZ)-1; added += 1
        seen.add(i)
        if str(g('מופיע כברירת מחדל') or '').strip() == 'כן': sel.append(i)
    removed = 0
    for i,r in enumerate(YZ):
        if i not in seen and r[3]: YZ[i] = [r[0], r[1], '', '', 0]; removed += 1   # נשמר המקום כדי שהמזהים לא יזוזו
    js = ('/* תאריכי פטירת גדולי ישראל ודברי ימי עם ישראל, מתוך לוח "עתים לבינה" תשפ"ז, באישור. [חודש, יום, שנה, טקסט, אישיות, מקור] */\n'
          'const YZ=[' + ','.join('[' + ','.join(esc(x) for x in r) + ']' for r in YZ) + '];\n'
          'const YZ_DEFAULT_SEL=' + json.dumps(sorted(set(sel))) + ';\n')
    SRC.write_text(js, encoding='utf-8')
    print('changed', changed, 'added', added, 'removed', removed)
    subprocess.run(['python3', str(ROOT/'tools/build_yz.py')], check=True)

if __name__ == '__main__':
    {'export': export, 'import': imp}[sys.argv[1] if len(sys.argv) > 1 else 'export']()
